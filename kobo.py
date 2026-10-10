"""Présence journalière des agents saisie avec KoboCollect.

Circuit :
  1. l'application génère le formulaire (XLSForm) avec la liste à jour des agents et des lieux (bureau, chantiers) ;
  2. le pointeur de chaque lieu fait l'appel sur son téléphone : une fiche par lieu et par jour, une ligne par agent ;
  3. l'application récupère les fiches (API KoboToolbox, ou fichier Excel exporté de Kobo), les contrôle avec les
     mêmes règles que le pointage manuel, puis les enregistre dans les pointages (table « attendance »).

Chaque ligne de fiche Kobo traitée est notée dans la table « kobo_pointages » : elle n'est jamais importée deux fois.
"""
from __future__ import annotations

import datetime as dt
import io
import re

import openpyxl
import streamlit as st
from openpyxl.styles import Font, PatternFill

import auth
import db
import regles as R

SERVEURS = {"https://kf.kobotoolbox.org": "KoboToolbox (serveur mondial)",
            "https://eu.kobotoolbox.org": "KoboToolbox (serveur européen)"}
STATUTS = {"present": "Présent", "absent": "Absent", "conge": "Congé", "maladie": "Maladie",
           "mission": "Mission", "repos": "Repos"}
AVEC_HEURES = {"present", "mission"}
FORM_ID = "presence_agents_mydestiny"

# États d'une ligne après contrôle
NOUVEAU, CONFLIT, IDENTIQUE, ERREUR = "Nouveau", "Déjà pointé (différent)", "Déjà pointé (identique)", "À corriger"


# ------------------------------------------------------------------ codes des choix du formulaire
def code_agent(matricule: str) -> str:
    return "m_" + re.sub(r"[^A-Za-z0-9_]", "_", matricule.strip())


def code_lieu(projet_id: str) -> str:
    return "c_" + re.sub(r"[^A-Za-z0-9_]", "_", projet_id)


# ------------------------------------------------------------------ réglages
def reglages() -> dict:
    """Serveur et identifiant du formulaire : réglages de l'application (ou secrets) ; jeton d'accès : secrets seulement."""
    r = db.parametre("kobo", {}) or {}
    try:
        sec = dict(st.secrets.get("kobo", {}))
    except Exception:  # noqa: BLE001 — pas de fichier de secrets en local
        sec = {}
    serveur = r.get("serveur") or sec.get("serveur") or "https://kf.kobotoolbox.org"
    return {"serveur": str(serveur).rstrip("/"), "jeton": str(sec.get("jeton") or auth.secret("KOBO_JETON") or ""),
            "formulaire": str(r.get("formulaire") or sec.get("formulaire") or "").strip()}


# ------------------------------------------------------------------ 1. formulaire XLSForm
def xlsform(s: dict) -> tuple[bytes, dict]:
    """Formulaire KoboCollect avec les agents actifs et les lieux actuels. Renvoie (fichier, résumé)."""
    agents = sorted([p for p in s["personnel"] if p["active"]], key=lambda p: p["name"].lower())
    codes = [code_agent(p["code"]) for p in agents]
    doublons = sorted({c for c in codes if codes.count(c) > 1})
    if doublons:
        raise ValueError("Des matricules ne se distinguent que par des espaces ou des signes ("
                         + ", ".join(doublons) + "). Rendez-les distincts dans « Personnel » avant de générer le formulaire.")
    if not agents:
        raise ValueError("Aucun agent actif : ajoutez le personnel avant de générer le formulaire.")
    chantiers = [p for p in s["projects"] if p["status"] != "Terminé"]

    survey = [
        ["type", "name", "label", "hint", "required", "relevant", "constraint", "constraint_message",
         "appearance", "default", "calculation"],
        ["start", "debut"], ["end", "fin"], ["username", "utilisateur"],
        ["date", "date_pointage", "Date du pointage", "", "yes", "", ". <= today()",
         "La date ne peut pas être dans le futur.", "", "today()"],
        ["select_one lieu", "lieu", "Lieu (bureau ou chantier)", "", "yes", "", "", "", "minimal autocomplete"],
        ["text", "pointeur", "Nom du pointeur", "La personne qui fait l'appel", "yes"],
        ["geopoint", "position", "Position GPS du lieu", "Facultatif"],
        ["begin_repeat", "agents", "Agent pointé"],
        ["select_one agent", "agent", "Agent", "Tapez le nom ou le matricule", "yes", "", "", "", "minimal autocomplete"],
        ["select_one statut", "statut", "Présence", "", "yes", "", "", "", "horizontal"],
        ["time", "arrivee", "Heure d'arrivée", "", "yes", "${statut}='present' or ${statut}='mission'"],
        ["time", "depart", "Heure de départ", "Pointez le départ en fin de journée (gardez la fiche en brouillon d'ici là).",
         "yes", "${statut}='present' or ${statut}='mission'", "decimal-time(.) > decimal-time(${arrivee})",
         "Le départ doit suivre l'arrivée."],
        ["integer", "pause", "Pause (minutes)", "Temps non travaillé dans la journée", "yes",
         "${statut}='present' or ${statut}='mission'", ". >= 0 and . <= 240", "Entre 0 et 240 minutes.", "", "60"],
        ["calculate", "heures", "", "", "", "", "", "", "", "",
         "if(${depart} != '' and ${arrivee} != '', round((decimal-time(${depart}) - decimal-time(${arrivee})) * 24 "
         "- ${pause} div 60, 2), 0)"],
        ["note", "heures_note", "Heures travaillées : ${heures} h", "", "", "${statut}='present' or ${statut}='mission'"],
        ["text", "observation", "Observation", "Facultatif"],
        ["end_repeat", "agents"],
    ]
    choices = [["list_name", "name", "label"]]
    choices += [["lieu", "bureau", "Bureau"]] + [["lieu", code_lieu(p["id"]), p["name"]] for p in chantiers]
    choices += [["statut", k, v] for k, v in STATUTS.items()]
    choices += [["agent", c, f"{p['name']} ({p['code']})" + (f" – {p['job']}" if p.get("job") else "")]
                for c, p in zip(codes, agents)]
    version = dt.datetime.now().strftime("%Y%m%d%H%M")
    settings = [["form_title", "form_id", "version"],
                [f"Présence journalière des agents – {R.company().get('name', 'MY DESTINY SARL')}", FORM_ID, version]]

    wb = openpyxl.Workbook()
    for i, (nom, lignes) in enumerate([("survey", survey), ("choices", choices), ("settings", settings)]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nom
        for ligne in lignes:
            ws.append(ligne)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E9EEF4")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(x.value or "")) for x in col)) + 2)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"agents": len(agents), "chantiers": len(chantiers), "version": version}


# ------------------------------------------------------------------ 2. récupération des fiches
def _court(cle: str) -> str:
    return cle.rsplit("/", 1)[-1]


def _heure(v) -> str:
    m = re.search(r"(\d{1,2}):(\d{2})", str(v or ""))
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else ""


def _nombre(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _lignes(fiche: dict, agents: list[dict]) -> list[dict]:
    """Une fiche Kobo (champs à plat) + ses lignes d'agents → lignes normalisées."""
    f = {_court(k): v for k, v in fiche.items()}
    uuid = str(f.get("_uuid") or f.get("instanceID", "")).replace("uuid:", "")
    out = []
    for n, a in enumerate(agents, start=1):
        a = {_court(k): v for k, v in a.items()}
        out.append({
            "cle": f"{uuid}#{n}", "fiche": uuid, "soumis_le": str(f.get("_submission_time") or "")[:19],
            "date": str(f.get("date_pointage") or "")[:10], "lieu": str(f.get("lieu") or ""),
            "pointeur": str(f.get("pointeur") or "").strip(), "agent": str(a.get("agent") or ""),
            "statut": str(a.get("statut") or ""), "arrivee": _heure(a.get("arrivee")), "depart": _heure(a.get("depart")),
            "pause": _nombre(a.get("pause")) or 0, "heures_form": _nombre(a.get("heures")),
            "observation": str(a.get("observation") or "").strip(),
        })
    return out


def depuis_api(reg: dict) -> list[dict]:
    import requests

    if not reg["jeton"] or not reg["formulaire"]:
        raise ValueError("Indiquez le jeton d'accès (secrets) et l'identifiant du formulaire (réglages).")
    url = f"{reg['serveur']}/api/v2/assets/{reg['formulaire']}/data/"
    params = {"format": "json", "limit": 1000}
    lignes = []
    while url:
        try:
            r = requests.get(url, params=params, headers={"Authorization": f"Token {reg['jeton']}"}, timeout=40)
        except requests.RequestException as e:
            raise ValueError(f"Serveur Kobo injoignable ({e.__class__.__name__}). Vérifiez la connexion et réessayez.")
        if r.status_code in (401, 403):
            raise ValueError("Kobo refuse l'accès : vérifiez le jeton d'accès et que le compte a accès au formulaire.")
        if r.status_code == 404:
            raise ValueError("Formulaire introuvable : vérifiez son identifiant (ex. aBcD12eFgH34iJkL) et le serveur.")
        if not r.ok:
            raise ValueError(f"Kobo a répondu par une erreur {r.status_code}. Réessayez dans quelques minutes.")
        data = r.json()
        for fiche in data.get("results", []):
            lignes += _lignes(fiche, fiche.get("agents") or next((v for k, v in fiche.items()
                                                                     if _court(k) == "agents"), []) or [])
        url, params = data.get("next"), None
    return lignes


def depuis_excel(contenu: bytes) -> list[dict]:
    """Export Excel de Kobo (« Valeurs et en-têtes XML ») : feuille principale + feuille « agents »."""
    wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    feuilles = {ws.title: [[c for c in r] for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    if len(feuilles) < 2:
        raise ValueError("Le fichier doit contenir deux feuilles : les fiches et la liste « agents ». Exportez depuis Kobo "
                         "au format XLS avec « Valeurs et en-têtes XML ».")
    noms = list(feuilles)

    def dicts(lignes):
        if not lignes:
            return []
        tete = [_court(str(c or "")) for c in lignes[0]]
        return [dict(zip(tete, r)) for r in lignes[1:] if any(v not in (None, "") for v in r)]

    fiches = dicts(feuilles[noms[0]])
    rep = dicts(feuilles.get("agents") or feuilles[noms[1]])
    if fiches and "lieu" not in fiches[0]:
        raise ValueError("Colonnes introuvables : exportez avec « Valeurs et en-têtes XML » (et non les libellés).")
    par_index = {}
    for a in rep:
        cle = a.get("_submission__uuid") or a.get("_parent_index")
        par_index.setdefault(str(cle), []).append(a)
    lignes = []
    for f in fiches:
        lignes += _lignes(f, par_index.get(str(f.get("_uuid")), []) or par_index.get(str(f.get("_index")), []))
    return lignes


# ------------------------------------------------------------------ 3. contrôle et import
def deja_traitees() -> set[str]:
    return {r["id"] for r in db.tout("kobo_pointages", sans=("donnees",))}


def analyser(lignes: list[dict], s: dict) -> list[dict]:
    """Contrôle chaque ligne avec les règles du pointage manuel ; ajoute l'état et le pointage à écrire."""
    agents = {code_agent(p["code"]): p for p in s["personnel"]}
    lieux = {"bureau": {"location": "Bureau", "project": ""}}
    lieux.update({code_lieu(p["id"]): {"location": "Chantier", "project": p["id"]} for p in s["projects"]})
    existants = {(a["person_id"], a["date"]): a for a in s["attendance"]}
    vus, faites, out = set(), deja_traitees(), []
    auj = R.today()
    for l in lignes:
        if l["cle"] in faites:
            continue
        p = agents.get(l["agent"])
        r = {**l, "person_id": p["id"] if p else "", "nom": p["name"] if p else l["agent"].removeprefix("m_"),
             "matricule": p["code"] if p else "", "statut_lib": STATUTS.get(l["statut"], l["statut"]),
             "heures": 0.0, "etat": ERREUR, "message": "", "pointage": None, "existant": None}
        lieu = lieux.get(l["lieu"])
        if l["statut"] in AVEC_HEURES:
            if l["arrivee"] and l["depart"]:
                h1, m1 = map(int, l["arrivee"].split(":"))
                h2, m2 = map(int, l["depart"].split(":"))
                r["heures"] = round(max(0.0, (h2 * 60 + m2 - h1 * 60 - m1 - (l["pause"] or 0)) / 60), 2)
            elif l["heures_form"] is not None:
                r["heures"] = round(l["heures_form"], 2)
        erreur = (
            "Agent inconnu : il n'est plus dans le personnel ou son matricule a changé." if not p else
            "Agent inactif dans le personnel." if not p["active"] else
            "Date du pointage manquante ou invalide." if not R.is_date(l["date"]) else
            "Date dans le futur." if l["date"] > auj else
            f"Date antérieure à l'engagement de l'agent ({p['start']})." if l["date"] < p["start"] else
            "Lieu inconnu (chantier supprimé ?)." if not lieu else
            "Présence non reconnue." if l["statut"] not in STATUTS else
            "Heures d'arrivée et de départ manquantes (fiche envoyée avant la fin de journée ?)."
            if l["statut"] in AVEC_HEURES and not (l["arrivee"] and l["depart"]) and l["heures_form"] is None else
            "Plus de 24 heures dans la journée." if r["heures"] > 24 else
            "Agent pointé deux fois pour cette date dans les fiches Kobo." if (p["id"], l["date"]) in vus else "")
        if erreur:
            r["message"] = erreur
            out.append(r)
            continue
        vus.add((p["id"], l["date"]))
        note = " – ".join(x for x in [f"KoboCollect, pointeur {l['pointeur']}" if l["pointeur"] else "KoboCollect",
                                      l["observation"]] if x)
        r["pointage"] = {"person_id": p["id"], "date": l["date"], "status": STATUTS[l["statut"]],
                         "hours": float(r["heures"]), "note": note[:1000], **lieu}
        ex = existants.get((p["id"], l["date"]))
        r["existant"] = ex
        if ex is None:
            r["etat"] = NOUVEAU
        elif ex["status"] == r["pointage"]["status"] and abs((ex["hours"] or 0) - r["heures"]) < 0.01 \
                and (ex.get("project") or "") == lieu["project"]:
            r["etat"] = IDENTIQUE
        else:
            r["etat"] = CONFLIT
            r["message"] = f"Déjà pointé : {ex['status']}, {ex['hours']:g} h, {R.hr_place_label(ex, s)}"
        out.append(r)
    return out


def _trace(t, r: dict, resultat: str, attendance_id=None) -> None:
    u = auth.utilisateur() or {}
    t.inserer("kobo_pointages", {
        "id": r["cle"], "fiche": r["fiche"], "soumis_le": r["soumis_le"], "date": r["date"],
        "person_id": r["person_id"], "agent": r["agent"], "statut": r["statut_lib"], "heures": float(r["heures"]),
        "lieu": r["lieu"], "pointeur": r["pointeur"], "resultat": resultat, "attendance_id": attendance_id,
        "donnees": {k: r[k] for k in ("arrivee", "depart", "pause", "observation", "message")},
        "traite_le": db.maintenant(), "traite_par": u.get("nom", "")}, resume=f"Ligne KoboCollect {resultat.lower()}")


def importer(analyse: list[dict], remplacer: bool) -> dict:
    """Écrit les nouveaux pointages (et remplace les différents si demandé), dans une seule transaction."""
    compte = {"Importé": 0, "Remplacé": 0, "Identique": 0}
    with db.transaction("Import des pointages KoboCollect") as t:
        for r in analyse:
            if r["etat"] == NOUVEAU:
                id_ = t.inserer("attendance", r["pointage"],
                                resume=f"Pointage KoboCollect de {r['nom']} le {r['date']}")
                _trace(t, r, "Importé", id_)
                compte["Importé"] += 1
            elif r["etat"] == IDENTIQUE:
                _trace(t, r, "Identique", r["existant"]["id"])
                compte["Identique"] += 1
            elif r["etat"] == CONFLIT and remplacer:
                t.maj("attendance", r["existant"]["id"], r["pointage"],
                      resume=f"Pointage de {r['nom']} le {r['date']} remplacé par KoboCollect")
                _trace(t, r, "Remplacé", r["existant"]["id"])
                compte["Remplacé"] += 1
    return compte


def ignorer(lignes: list[dict], resultat: str = "Ignoré") -> int:
    with db.transaction("Lignes KoboCollect écartées") as t:
        for r in lignes:
            _trace(t, r, resultat, (r.get("existant") or {}).get("id"))
    return len(lignes)


# ------------------------------------------------------------------ 4. présence d'une journée
def presence_du_jour(s: dict, date: str) -> dict:
    """Pointages d'une date (manuels et KoboCollect) et agents actifs non pointés."""
    noms = {p["id"]: p for p in s["personnel"]}
    jour = [a for a in s["attendance"] if a["date"] == date and a["person_id"] in noms]
    pointes = {a["person_id"] for a in jour}
    actifs = [p for p in s["personnel"] if p["active"] and p["start"] <= date]
    manquants = sorted([p for p in actifs if p["id"] not in pointes], key=lambda p: p["name"].lower())
    par_lieu: dict[str, dict] = {}
    for a in jour:
        lib = R.hr_place_label(a, s)
        d = par_lieu.setdefault(lib, {"lieu": lib, "presents": 0, "absents": 0, "autres": 0, "heures": 0.0})
        d["presents" if a["status"] in ("Présent", "Mission") else "absents" if a["status"] == "Absent" else "autres"] += 1
        d["heures"] += a["hours"] or 0
    return {"jour": jour, "manquants": manquants, "par_lieu": sorted(par_lieu.values(), key=lambda d: d["lieu"]),
            "presents": sum(1 for a in jour if a["status"] in ("Présent", "Mission")),
            "absents": sum(1 for a in jour if a["status"] == "Absent"),
            "autres": sum(1 for a in jour if a["status"] not in ("Présent", "Mission", "Absent")),
            "heures": sum(a["hours"] or 0 for a in jour),
            "kobo": sum(1 for a in jour if (a.get("note") or "").startswith("KoboCollect"))}
