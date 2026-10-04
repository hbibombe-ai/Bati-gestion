"""Reprise historique : modèle Excel normalisé (un chantier par classeur) et contrôles avant conservation.

Un lot importé reste une donnée historique : il ne crée aucune écriture comptable et aucune dépense
n'est considérée comme justifiée par le seul import.
"""
from __future__ import annotations

import datetime as dt
import io
import re

import openpyxl
from openpyxl.styles import Font, PatternFill

SCHEMA = {
    "Chantier": ["reference", "nom", "client_ref", "lieu", "contrat", "date_debut", "date_fin", "monnaie", "montant"],
    "Clients": ["reference", "nom", "telephone", "email", "adresse"],
    "Fournisseurs": ["reference", "nom", "telephone", "email", "adresse"],
    "DQE_BOQ": ["reference", "description", "unite", "quantite_prevue", "quantite_executee", "prix_unitaire", "monnaie"],
    "Budgets": ["reference", "description", "date", "montant", "monnaie"],
    "Devis": ["reference", "description", "client_ref", "date", "montant", "monnaie"],
    "Depenses": ["reference", "description", "fournisseur_ref", "responsable_ref", "date", "montant", "monnaie"],
    "Paiements": ["reference", "depense_ref", "date", "montant", "monnaie", "compte"],
    "Stocks": ["reference", "description", "unite", "quantite"],
    "Materiels": ["reference", "description", "numero_serie", "etat"],
    "Personnel": ["reference", "nom", "fonction", "date_debut", "date_fin"],
    "Affectations": ["reference", "personnel_ref", "fonction", "date_debut", "date_fin"],
    "Situations": ["reference", "description", "date", "montant", "monnaie", "avancement"],
    "Documents": ["reference", "depense_ref", "nom_fichier", "type", "date", "montant", "monnaie"],
}
MONEY = {"montant", "prix_unitaire"}
NUMERIC = MONEY | {"quantite", "quantite_prevue", "quantite_executee", "avancement"}
REFS = {"client_ref": "Clients", "fournisseur_ref": "Fournisseurs", "responsable_ref": "Personnel",
        "personnel_ref": "Personnel", "depense_ref": "Depenses"}


def _date_ok(v: str) -> bool:
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", v or ""):
        return False
    try:
        return dt.date.fromisoformat(v).isoformat() == v
    except ValueError:
        return False


def valider(tables: dict, refs_existantes: set[str]) -> dict:
    errors, warnings, rows = [], [], []

    def issue(sheet, line, field, message):
        errors.append({"sheet": sheet, "line": line, "field": field, "message": message})
    refs: dict[str, dict] = {}
    if len(tables.get("Chantier") or []) != 1:
        issue("Chantier", 2, "reference", "Un seul chantier est requis par fichier.")
    for sheet in tables:
        if sheet not in SCHEMA and sheet != "Instructions":
            issue(sheet, 1, "", "Feuille inconnue : utiliser le modèle normalisé.")
    for sheet, columns in SCHEMA.items():
        refs[sheet] = {}
        for i, source in enumerate(tables.get(sheet) or []):
            line = source.get("__line", i + 2)
            row = {}
            for key in columns:
                value = source.get(key, "")
                if value is None:
                    value = ""
                if isinstance(value, (dt.datetime, dt.date)):
                    value = (value.date() if isinstance(value, dt.datetime) else value).isoformat()
                if key in NUMERIC and value != "":
                    raw = str(value).strip().replace(",", ".")
                    ok = bool(re.fullmatch(r"\d+(\.\d+)?", raw)) and float(raw) <= 1e10
                    decimales = raw.split(".")[1].rstrip("0") if "." in raw else ""
                    if not ok or (key in MONEY and len(decimales) > 2):
                        issue(sheet, line, key, "Nombre positif requis ; montants avec deux décimales maximum.")
                        row[key] = 0
                    else:
                        row[key] = float(raw) if "." in raw else int(raw)
                else:
                    row[key] = str(value).strip()
                    if len(row[key]) > 2000:
                        issue(sheet, line, key, "Texte trop long (2 000 caractères maximum).")
            if not row.get("reference"):
                issue(sheet, line, "reference", "Référence obligatoire.")
            if row.get("reference") in refs[sheet]:
                issue(sheet, line, "reference", "Référence en double dans cette feuille.")
            refs[sheet][row.get("reference")] = row
            if "nom" in columns and not row["nom"]:
                issue(sheet, line, "nom", "Nom obligatoire.")
            if "description" in columns and not row["description"]:
                issue(sheet, line, "description", "Description obligatoire.")
            if "monnaie" in columns and row["monnaie"] not in ("USD", "CDF", "EUR"):
                issue(sheet, line, "monnaie", "Choisir USD, CDF ou EUR.")
            for key in [c for c in columns if c.startswith("date")]:
                if (key == "date" and not row[key]) or (row[key] and not _date_ok(row[key])):
                    issue(sheet, line, key, "Date requise au format AAAA-MM-JJ.")
            if row.get("date_debut") and row.get("date_fin") and row["date_fin"] < row["date_debut"]:
                issue(sheet, line, "date_fin", "La fin précède le début.")
            for key in [c for c in columns if c in NUMERIC]:
                if row[key] == "":
                    issue(sheet, line, key, "Valeur numérique obligatoire (0 si nul).")
            if isinstance(row.get("avancement"), (int, float)) and row["avancement"] > 100:
                issue(sheet, line, "avancement", "L’avancement ne peut dépasser 100 %.")
            if sheet == "Paiements" and row["compte"] not in ("cash", "bank", "mobile", "unassigned"):
                issue(sheet, line, "compte", "Choisir cash, bank, mobile ou unassigned.")
            rows.append({"sheet": sheet, "line": line, "values": row})
    if len(rows) > 5000:
        issue("Classeur", 0, "", "Maximum 5 000 lignes par lot.")
    for r in rows:
        sheet, line, row = r["sheet"], r["line"], r["values"]
        for field, target in REFS.items():
            if row.get(field) and row[field] not in refs[target]:
                issue(sheet, line, field, f"Référence absente de la feuille {target}.")
        if sheet == "Paiements" and not row.get("depense_ref"):
            issue(sheet, line, "depense_ref", "Relier le paiement à une dépense.")
        if sheet == "Affectations" and not row.get("personnel_ref"):
            issue(sheet, line, "personnel_ref", "Relier l’affectation à une personne.")
        if sheet == "Documents" and not row.get("nom_fichier"):
            issue(sheet, line, "nom_fichier", "Nom du fichier obligatoire.")
        if sheet == "Depenses":
            warnings.append({"sheet": sheet, "line": line, "field": "justificatifs",
                             "message": "Dépense historique non justifiée avant validation des pièces par Finance et DG."})
    paid: dict[str, int] = {}
    for r in [x for x in rows if x["sheet"] == "Paiements"]:
        row, line = r["values"], r["line"]
        dep = refs["Depenses"].get(row["depense_ref"])
        if dep and row["monnaie"] != dep["monnaie"]:
            issue("Paiements", line, "monnaie", "Le paiement et la dépense doivent avoir la même monnaie dans ce modèle.")
        total = paid.get(row["depense_ref"], 0) + round(float(row["montant"] or 0) * 100)
        paid[row["depense_ref"]] = total
        if dep and total > round(float(dep["montant"] or 0) * 100):
            issue("Paiements", line, "montant", "Les paiements cumulés dépassent la dépense.")
    projet = next((r["values"] for r in rows if r["sheet"] == "Chantier"), None)
    if projet and projet["reference"] in refs_existantes:
        issue("Chantier", 2, "reference", "Ce chantier possède déjà un lot préparé. Supprimez le lot précédent avant de "
                                          "le remplacer.")
    return {"valid": not errors, "errors": errors, "warnings": warnings, "rows": rows, "project": projet}


def lire_classeur(contenu: bytes) -> tuple[dict, list]:
    if len(contenu) > 8 * 1024 * 1024:
        raise ValueError("Utilisez un fichier .xlsx de 8 Mo maximum.")
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(contenu), data_only=False, read_only=False)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Fichier Excel illisible : utilisez le modèle .xlsx.") from e
    tables, errors = {}, []
    for ws in wb_f.worksheets:
        name = ws.title
        if name == "Instructions":
            continue
        if ws.max_row > 5001 or ws.max_column > 31:
            errors.append({"sheet": name, "line": 1, "field": "", "message": "Feuille trop grande : maximum 5 000 lignes "
                                                                            "et 31 colonnes."})
            continue
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    errors.append({"sheet": name, "line": cell.row, "field": cell.coordinate,
                                   "message": "Remplacer les formules par leurs valeurs avant import."})
        if name not in SCHEMA:
            tables[name] = []
            continue
        headers = [str(v).strip() if v is not None else "" for v in (grid[0] if grid else [])]
        while headers and headers[-1] == "":
            headers.pop()
        if headers != SCHEMA[name]:
            errors.append({"sheet": name, "line": 1, "field": "", "message": "Colonnes différentes du modèle. "
                                                                           "Télécharger le modèle normalisé."})
        lignes = []
        for i, values in enumerate(grid[1:], start=2):
            if all(v in ("", None) for v in values):
                continue
            lignes.append({"__line": i, **{h: (values[j] if j < len(values) else "") for j, h in enumerate(headers)}})
        tables[name] = lignes
    return tables, errors


def modele() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    gras = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="112740")
    for name, headers in SCHEMA.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for j, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=j)
            c.font, c.fill = gras, fond
            ws.column_dimensions[c.column_letter].width = max(18, len(h) + 3)
    ws = wb.create_sheet("Instructions")
    for l in [["MY DESTINY SARL — Reprise historique"], ["Version du modèle", "1"],
              ["Un seul chantier par fichier ; références uniques dans chaque feuille."],
              ["Conserver les en-têtes. Laisser les feuilles non utilisées vides."],
              ["Dates : AAAA-MM-JJ. Monnaies : USD, CDF, EUR. Montants : nombres, deux décimales maximum."],
              ["Les références client, fournisseur, personnel et dépense doivent exister dans ce même classeur."],
              ["Paiements : compte = cash, bank, mobile ou unassigned."],
              ["Documents : indiquer le nom du fichier, puis joindre le fichier au lot après le contrôle."],
              ["Les fichiers ne sont pas incorporés dans ce modèle Excel."],
              ["Aucune dépense ne sera considérée comme justifiée par le seul import."]]:
        ws.append(l)
    ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rapport_controle(res: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle"
    ws.append(["Niveau", "Feuille", "Ligne", "Champ", "Message"])
    for e in res["errors"]:
        ws.append(["Erreur", e["sheet"], e["line"], e["field"], e["message"]])
    for w in res["warnings"]:
        ws.append(["Avertissement", w["sheet"], w["line"], w["field"], w["message"]])
    if ws.max_row == 1:
        ws.append(["Information", "", "", "", "Aucune anomalie de structure détectée."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
