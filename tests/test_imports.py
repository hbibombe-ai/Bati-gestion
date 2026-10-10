"""Recette de l'import généralisé (spécification du 10 octobre 2026, §15), sur base SQLite isolée."""
import io
import os
import threading
import unittest
from unittest.mock import patch

import openpyxl
from sqlalchemy import create_engine, select
from sqlalchemy.pool import StaticPool

import auth
import db
import gestion_v7 as V
import migration
from imports import excel as X
from imports import service as SV

D = "2026-11-01"


def classeur(fonction, mode, ref, feuilles, date=D, source="Ancien tableur"):
    wb = openpyxl.load_workbook(io.BytesIO(X.modele(fonction, mode, date, ref)))
    p = wb["Parametres"]
    for row in p.iter_rows(min_row=2):
        if row[0].value == "source_systeme":
            row[1].value = source
    for nom, lignes in feuilles.items():
        ws = wb[nom]
        cols = [c.value for c in ws[1]]
        for i, l in enumerate(lignes, start=2):
            for j, c in enumerate(cols, start=1):
                if c in l:
                    ws.cell(row=i, column=j, value=l[c])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PG = os.environ.get("BG_TEST_PG")   # ex. postgresql+psycopg2://postgres@/bati_test?host=/var/tmp&port=55432


class Base(unittest.TestCase):
    def setUp(self):
        if PG:
            self.engine = create_engine(PG, pool_size=10)
            db.meta.drop_all(self.engine)
        else:
            self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        db.meta.create_all(self.engine)
        self.p1 = patch.object(db, "moteur", lambda: self.engine)
        self.p1.start()
        self.user = None
        self.p2 = patch.object(auth, "utilisateur", lambda: self.user)
        self.p2.start()
        roles = ["chantier", "finance", "finance", "admin", "admin", "finance", "rh", "chantier"]
        with self.engine.begin() as c:
            for i, role in enumerate(roles, 1):
                c.execute(db.users.insert().values(id=i, identifiant=str(i), nom=f"User {i}", role=role, mdp_hash="x",
                                                   actif=True, doit_changer_mdp=False))
            c.execute(db.projects.insert().values(id="p", name="Résidence Kintambo", currency="USD", status="En cours",
                                                  client="", budget=0, progress=10))
            c.execute(db.clients.insert().values(id="f", code="F001", name="Quincaillerie Lemba", roles=["supplier"]))

    def tearDown(self):
        self.p2.stop()
        self.p1.stop()
        self.engine.dispose()

    def as_(self, i):
        self.user = db.un("users", i)

    def importer(self, fonction, mode, contenu, uid=2, decisions=None, fichiers=None, date=D):
        self.as_(uid)
        ds = SV.parse_file(contenu, fonction)
        ctx = SV.contexte(mode, date if mode == "reprise" else "", fichiers=fichiers)
        pv = SV.preview_import(ds, ctx, contenu, "test.xlsx")
        if pv["preview_id"] is None or pv["rapport"]["errors"]:
            return pv
        res = SV.commit_import(pv["preview_id"], decisions or {"accepter_alertes": True}, fichiers)
        return {**pv, **res}

    def count(self, table):
        with self.engine.connect() as c:
            return len(c.execute(select(db.TABLES[table])).fetchall())

    def budget_approuve(self, montant=1000000):
        self.as_(1)
        b = V.create_budget("p", "CT-KIN", [dict(code="MAT", label="Matériaux", budget=montant, forecast=montant)],
                            montant, montant)
        self.as_(4)
        V.approve_budget(b, "ok")
        return b

    def compte(self, ref="BANQUE-USD", solde=0, uid=6):
        self.as_(uid)
        a = V.create_account(ref, "Banque", "USD", "5211", "bank")
        if solde:
            V.opening(a, solde, "2026-10-01", "init")
        return a

    def activer(self, lot_id, finance=3, dg=5):
        self.as_(finance)
        SV.valider_finance(lot_id, "Rapprochement relevé n°12")
        self.as_(dg)
        return SV.activer(lot_id, "Reprise validée")


class Referentiels(Base):
    def fichier_ref(self):
        return classeur("globale", "en_cours", "LOT-REF-1", {
            "Tiers": [{"reference_externe": "C010", "nom": "SCI Gombe", "roles": "client"},
                      {"reference_externe": "F001", "nom": "Quincaillerie Lemba", "roles": "fournisseur"},
                      {"reference_externe": "00123", "nom": "Transport Ndjili", "roles": "fournisseur ; sous-traitant"}],
            "Chantiers": [{"reference_externe": "CH-GOMBE", "nom": "Immeuble Gombe", "client_ref": "C010", "devise": "USD",
                           "etat": "En cours", "avancement": 35, "date_debut": "2026-01-10", "date_fin": "2027-06-30"}],
            "Personnel": [{"reference_externe": "M001", "nom": "Kabeya Jean", "fonction": "Maçon", "date_engagement": "2026-02-01",
                           "lieu": "Chantier", "chantier_ref": "CH-GOMBE", "actif": "oui"}],
            "Presences": [{"matricule": "M001", "date": "2026-10-05", "statut": "Présent", "heures": 8},
                          {"matricule": "M001", "date": "2026-10-06", "statut": "Absent", "heures": 0}]})

    def test_fichier_valide_relations_audit_et_idempotence(self):
        r = self.importer("globale", "en_cours", self.fichier_ref(), uid=4)
        self.assertFalse(r["existant"])
        self.assertEqual(r["lot"]["statut"], "Importé")
        clients = db.tout("clients")
        self.assertEqual(len(clients), 3)
        self.assertIn("00123", {c["code"] for c in clients})          # zéros initiaux conservés
        projet = next(p for p in db.tout("projects") if p["name"] == "Immeuble Gombe")
        self.assertEqual(projet["client"], next(c["id"] for c in clients if c["code"] == "C010"))
        pers = db.tout("personnel")[0]
        self.assertEqual(pers["project"], projet["id"])
        self.assertEqual(self.count("attendance"), 2)
        maps = {(m["entite"], m["reference_externe"]): m["effet"] for m in db.tout("import_mappings")}
        self.assertEqual(maps[("tiers", "F001")], "Correspondance")
        self.assertEqual(maps[("chantier", "CH-GOMBE")], "Création")
        self.assertTrue(any(a["objet"] == "import_lots" for a in db.tout("audit")))
        n_audit = self.count("audit")
        r2 = self.importer("globale", "en_cours", self.fichier_ref(), uid=4)           # même fichier, double clic
        self.assertTrue(r2.get("existant") or r2["rapport"]["errors"])
        self.assertEqual(len(db.tout("clients")), 3)
        self.assertEqual(self.count("attendance"), 2)
        self.assertEqual(self.count("import_lots"), 1)
        self.assertLessEqual(self.count("audit") - n_audit, 1)

    def test_double_confirmation_meme_apercu_et_apercus_concurrents(self):
        contenu = self.fichier_ref()
        self.as_(4)
        ds = SV.parse_file(contenu, "globale")
        a = SV.preview_import(ds, SV.contexte("en_cours"), contenu)
        b = SV.preview_import(ds, SV.contexte("en_cours"), contenu)
        r1 = SV.commit_import(a["preview_id"], {"accepter_alertes": True})
        r2 = SV.commit_import(a["preview_id"], {"accepter_alertes": True})
        r3 = SV.commit_import(b["preview_id"], {"accepter_alertes": True})
        self.assertFalse(r1["existant"])
        self.assertTrue(r2["existant"] and r3["existant"])
        self.assertEqual(self.count("import_lots"), 1)
        self.assertEqual(self.count("attendance"), 2)
        with self.assertRaises(Exception):
            with db.transaction("doublon") as t:
                t.inserer("import_mappings", {"id": "x", "lot_id": "l", "entite": "tiers", "reference_externe": "C010",
                                              "id_interne": "y", "effet": "Création"})

    def test_erreur_derniere_ligne_aucune_ecriture_partielle(self):
        lignes = [{"reference_externe": f"T{i}", "nom": f"Tiers {i}", "roles": "client"} for i in range(5)]
        lignes[-1]["email"] = "pas-une-adresse"
        r = self.importer("tiers", "en_cours", classeur("tiers", "en_cours", "LOT-ERR", {"Tiers": lignes}), uid=2)
        self.assertTrue(r["rapport"]["errors"])
        self.assertEqual(self.count("clients"), 1)
        # Échec à l'exécution de la dernière action : tout le lot est annulé, y compris le lot et les correspondances.
        lignes[-1].pop("email")
        contenu = classeur("tiers", "en_cours", "LOT-ERR2", {"Tiers": lignes})
        self.as_(2)
        pv = SV.preview_import(SV.parse_file(contenu, "tiers"), SV.contexte("en_cours"), contenu)
        from imports.adapters import referentiels as RF
        original = RF.executer_tiers
        def casse(t, ctx, a, res):
            if a["ref"] == "T4":
                raise ValueError("panne simulée")
            return original(t, ctx, a, res)
        with patch.dict(RF.TIERS, {"executer": casse}):
            with self.assertRaises(ValueError):
                SV.commit_import(pv["preview_id"], {})
        self.assertEqual((self.count("clients"), self.count("import_lots"), self.count("import_mappings")), (1, 0, 0))

    def test_presences_doublons_kobo_et_conflits(self):
        self.importer("globale", "en_cours", self.fichier_ref(), uid=4)
        pid = db.tout("personnel")[0]["id"]
        contenu = classeur("presences", "en_cours", "LOT-PRES", {"Presences": [
            {"matricule": "M001", "date": "2026-10-05", "statut": "Présent", "heures": 8},     # identique (ex. Kobo)
            {"matricule": "M001", "date": "2026-10-06", "statut": "Présent", "heures": 8}]})  # conflit
        r = self.importer("presences", "en_cours", contenu, uid=7)
        self.assertTrue(any("Conflit" in e["message"] for e in r["rapport"]["errors"]))
        contenu = classeur("presences", "en_cours", "LOT-PRES2", {"Presences": [
            {"matricule": "M001", "date": "2026-10-05", "statut": "Présent", "heures": 8},
            {"matricule": "M001", "date": "2026-10-07", "statut": "Présent", "heures": 7.5}]})
        r = self.importer("presences", "en_cours", contenu, uid=7)
        self.assertEqual(r["rapport"]["comptes"].get("Ignoré (identique)"), 1)
        self.assertEqual(len([a for a in db.tout("attendance") if a["person_id"] == pid]), 3)


class Securite(Base):
    def test_formules_macros_nombres_ambigus_et_dates(self):
        wb = openpyxl.load_workbook(io.BytesIO(X.modele("tiers", "en_cours", "", "L")))
        wb["Tiers"]["A2"] = "T1"
        wb["Tiers"]["B2"] = "=1+1"
        wb["Tiers"]["C2"] = "client"
        buf = io.BytesIO()
        wb.save(buf)
        self.as_(2)
        ds = SV.parse_file(buf.getvalue(), "tiers")
        self.assertTrue(any("Formule" in e["message"] for e in ds["errors"]))
        with self.assertRaises(ValueError):
            SV.parse_file(b"PK\x03\x04" + b"0" * 100, "tiers")
        for bad in ("1.234,56", "1 234,56", "1,234.56", "12.345"):
            with self.assertRaises(ValueError):
                X.centimes(bad)
        self.assertEqual(X.centimes("1234,56"), 123456)
        self.assertEqual(X.centimes(1000.1), 100010)
        champ = X.S.Champ("d", "date", True)
        with self.assertRaises(ValueError):
            X.normaliser(champ, "09/10/2026")
        self.assertEqual(X.normaliser(X.S.Champ("p", "month"), "2026-10"), "2026-10")

    def test_feuille_inconnue_et_colonne_statut_approuve(self):
        wb = openpyxl.load_workbook(io.BytesIO(X.modele("tiers", "en_cours", "", "L")))
        wb.create_sheet("Divers")
        wb["Tiers"].cell(row=1, column=8, value="statut_approuve")
        buf = io.BytesIO()
        wb.save(buf)
        self.as_(2)
        ds = SV.parse_file(buf.getvalue(), "tiers")
        msgs = " ".join(e["message"] for e in ds["errors"])
        self.assertIn("Feuille inconnue", msgs)
        self.assertIn("approuvé", msgs)

    def test_droits_et_donnees_rh_confidentielles(self):
        contenu = classeur("personnel", "en_cours", "LOT-RH", {"Personnel": [
            {"reference_externe": "M100", "nom": "Mbuyi Paul", "fonction": "Chef d’équipe", "date_engagement": "2026-01-05",
             "lieu": "Bureau", "actif": "oui", "salaire_base": "450", "devise_salaire": "USD", "periode_salaire": "Mois",
             "salaire_effet": "2026-01-05"}]})
        r = self.importer("personnel", "en_cours", contenu, uid=1)        # conducteur de travaux
        self.assertIsNone(r["preview_id"])
        self.assertTrue(r["rapport"]["confidentiel_refuse"])
        self.assertEqual(r["rapport"]["plan"], [])
        self.assertEqual(self.count("import_previews"), 0)
        r = self.importer("personnel", "en_cours", contenu, uid=2)        # finance : pas de droit RH
        self.assertIsNone(r["preview_id"])
        r = self.importer("personnel", "en_cours", contenu, uid=7)        # RH
        self.assertEqual(db.tout("personnel")[0]["salary"]["amount"], 45000)
        masque = SV.lignes_lot(r["lot"]["id"], "finance")["Personnel"][0]
        self.assertEqual(masque["salaire_base"], "<confidentiel>")


class Reprise(Base):
    def test_archives_seules_sans_effet(self):
        a = self.compte(solde=500000)
        self.budget_approuve()
        avant = (self.count("movements"), len(V.records("entry")), V.records("budget")[0]["data"], V.balance(db.un("v7_records", a)))
        contenu = classeur("globale", "archives", "ARCH-2025", {
            "Factures_Fournisseurs": [{"reference_externe": "FA1", "numero_facture": "F-2025-77", "tiers_ref": "F001",
                                       "chantier_ref": "p", "date_facture": "2025-03-01", "date_echeance": "2025-04-01",
                                       "devise": "USD", "montant": 1000}],
            "Paiements": [{"reference_externe": "PA1", "facture_ref": "FA1", "date": "2025-03-15", "devise": "USD", "montant": 600}],
            "Mouvements_Tresorerie": [{"reference_externe": "MV1", "compte_ref": "BANQUE-USD", "date": "2025-03-15",
                                       "sens": "sortie", "devise": "USD", "montant": 600, "libelle": "Paiement F-2025-77"}],
            "Paie": [], "Stocks": [{"reference_externe": "S1", "article": "Ciment", "quantite": 40}]}, date="")
        r = self.importer("globale", "archives", contenu, uid=4)
        self.assertEqual(r["lot"]["statut"], "Archivé")
        apres = (self.count("movements"), len(V.records("entry")), V.records("budget")[0]["data"], V.balance(db.un("v7_records", a)))
        self.assertEqual(avant, apres)
        self.assertEqual(V.records("carryover"), [])
        self.assertEqual(len(SV.lignes_lot(r["lot"]["id"], "admin")["Factures_Fournisseurs"]), 1)

    def test_facture_1000_payee_600_reliquat_400_sans_redebit(self):
        self.budget_approuve()
        a = self.compte(solde=200000)
        contenu = classeur("fournisseurs", "reprise", "REPRISE-FOUR", {
            "Factures_Fournisseurs": [{"reference_externe": "FA1", "numero_facture": "F-2026-15", "tiers_ref": "F001",
                                       "chantier_ref": "p", "code_dqe": "MAT", "date_facture": "2026-09-01",
                                       "date_echeance": "2026-11-30", "devise": "USD", "montant": "1000"}],
            "Paiements": [{"reference_externe": "PA1", "facture_ref": "FA1", "date": "2026-09-20", "devise": "USD",
                           "montant": "600", "compte": "bank"}]})
        r = self.importer("fournisseurs", "reprise", contenu, uid=2)
        self.assertEqual(r["lot"]["statut"], "Préparé")
        c = V.records("carryover")[0]
        self.assertEqual((c["data"]["type"], c["data"]["amount"], c["data"]["status"]), ("debt", 40000, "À approuver"))
        self.assertEqual(c["data"]["archive"]["paye_avant_bascule"], 60000)
        solde = V.balance(db.un("v7_records", a))
        self.as_(2)
        with self.assertRaises(ValueError):                    # l'importateur ne valide pas son propre lot
            SV.valider_finance(r["lot"]["id"], "moi")
        self.activer(r["lot"]["id"])
        self.assertEqual(V.balance(db.un("v7_records", a)), solde)       # ancien paiement non redébité
        self.as_(5)
        with self.assertRaises(ValueError):                    # la DG qui a activé ne paie pas
            V.pay_carryover(c["id"], a, 40000, "2026-11-05", "VIR-1")
        self.as_(6)
        with self.assertRaises(ValueError):
            V.pay_carryover(c["id"], a, 40000, "2026-10-20", "VIR-0")   # antérieur à la bascule : archives
        V.pay_carryover(c["id"], a, 40000, "2026-11-05", "VIR-1")
        with self.assertRaises(ValueError):
            V.pay_carryover(c["id"], a, 1, "2026-11-06", "VIR-2")
        self.assertEqual(db.un("v7_records", c["id"])["data"]["status"], "Soldée")
        self.assertEqual(V.balance(db.un("v7_records", a)), solde - 40000)
        self.assertEqual(sum(x["debit"] - x["credit"] for x in V.ledger()), 0)
        m = V.metrics(V.records("budget")[0], V.records("operation"))
        self.assertEqual((m["cost"], m["cash"]), (0, 40000))   # la dette n'ajoute pas un second coût

    def test_avance_500_justifiee_350_reliquat_150(self):
        self.budget_approuve()
        a = self.compte(solde=100000)
        contenu = classeur("fournisseurs", "reprise", "REPRISE-AV", {
            "Avances": [{"reference_externe": "AV1", "beneficiaire_ref": "F001", "chantier_ref": "p", "code_dqe": "MAT",
                         "objet": "Régie chantier", "date_versement": "2026-09-01", "devise": "USD", "montant_verse": 500,
                         "montant_justifie_declare": 350, "montant_restitue": 0}]})
        r = self.importer("fournisseurs", "reprise", contenu, uid=2)
        self.assertTrue(any(w["decision"] for w in r["rapport"]["warnings"]))     # justification déclarée ≠ preuve
        mouvements = self.count("movements")
        self.activer(r["lot"]["id"])
        c = V.records("carryover")[0]
        self.assertEqual(c["data"]["amount"], 15000)
        self.assertEqual(c["data"]["archive"]["montant_verse"], 50000)
        self.assertEqual(self.count("movements"), mouvements)             # aucun versement fictif
        m = V.metrics(V.records("budget")[0], V.records("operation"))
        self.assertEqual((m["advance"], m["cost"]), (15000, 0))
        self.as_(1)
        V.upload_carryover_piece(c["id"], "facture.pdf", b"%PDF-reprise-test", "2026-11-03", 10000)
        piece = db.un("v7_records", c["id"])["data"]["justifications"][0]["piece"]
        self.as_(3)
        V.validate_carryover_justification(c["id"], piece, True, "Pièce contrôlée")
        self.as_(6)
        V.return_carryover_advance(c["id"], a, 5000, "2026-11-04", "RESTIT-1")
        m = V.metrics(V.records("budget")[0], V.records("operation"))
        self.assertEqual((m["advance"], m["cost"]), (0, 10000))
        self.assertEqual(db.un("v7_records", c["id"])["data"]["status"], "Soldée")

    def test_budget_10000_couts_3000_engagements_2000_disponible_5000(self):
        contenu = classeur("budgets", "reprise", "REPRISE-BUD", {
            "Budgets": [{"reference_externe": "CT-KIN", "chantier_ref": "p", "devise": "USD", "montant_contrat": 15000,
                         "cout_estime_terminaison": 10000}],
            "Lignes_Budget": [{"budget_ref": "CT-KIN", "numero_ligne": 1, "code_dqe": "MAT", "libelle": "Matériaux",
                               "budget_initial": 10000, "quantite": 100, "prix_unitaire": 100}],
            "Cumuls_Budget": [{"reference_externe": "CU1", "chantier_ref": "p", "code_dqe": "MAT", "devise": "USD",
                               "couts_executes": 3000, "engagements_ouverts": 2000, "date_arrete": "2026-10-31"}]})
        r = self.importer("budgets", "reprise", contenu, uid=2)
        self.assertEqual(r["lot"]["statut"], "Préparé")
        self.assertEqual(V.records("budget")[0]["data"]["status"], "À approuver")
        self.activer(r["lot"]["id"])
        b = V.records("budget")[0]
        self.assertEqual(b["data"]["status"], "Approuvé")
        m = V.metrics(b, V.records("operation"))
        self.assertEqual((m["budget"], m["cost"], m["open"], m["available"], m["cash"]), (1000000, 300000, 200000, 500000, 0))
        self.assertEqual(V.records("payment"), [])
        rap = SV.lot(r["lot"]["id"])["rapprochement"]
        self.assertTrue(any(e["rubrique"] == "Budget Disponible" for e in rap["ecarts_activation"]))
        engagement = next(c for c in V.records("carryover") if c["data"]["type"] == "commitment")
        self.as_(3)
        V.close_commitment(engagement["id"], "Remplacé par la demande N0 OP-77")
        self.assertEqual(V.metrics(V.records("budget")[0], V.records("operation"))["available"], 700000)

    def test_tresorerie_et_balance_sans_double_comptabilisation(self):
        double = classeur("globale", "reprise", "REPRISE-TRESO", {
            "Comptes_Tresorerie": [{"reference_externe": "CAISSE-USD", "libelle": "Caisse siège", "devise": "USD",
                                    "compte_comptable": "5711", "rattachement": "cash", "solde_ouverture": 250,
                                    "date_solde": D, "reference_releve": "PV comptage 31/10"}],
            "Ecritures": [{"reference_externe": "BO-USD", "date": D, "journal": "OD", "libelle": "Balance d’ouverture", "devise": "USD"}],
            "Lignes_Ecriture": [{"ecriture_ref": "BO-USD", "numero_ligne": 1, "compte": "5711", "debit": 250, "credit": 0},
                                {"ecriture_ref": "BO-USD", "numero_ligne": 2, "compte": "101", "debit": 0, "credit": 250}]})
        r = self.importer("globale", "reprise", double, uid=2)
        self.assertTrue(any("double comptabilisation" in e["message"] for e in r["rapport"]["errors"]))
        desequilibre = classeur("comptabilite", "reprise", "REPRISE-BO", {
            "Ecritures": [{"reference_externe": "BO", "date": D, "journal": "OD", "libelle": "Ouverture", "devise": "CDF"}],
            "Lignes_Ecriture": [{"ecriture_ref": "BO", "numero_ligne": 1, "compte": "4011", "debit": 0, "credit": 500},
                                {"ecriture_ref": "BO", "numero_ligne": 2, "compte": "101", "debit": 400, "credit": 0}]})
        r = self.importer("comptabilite", "reprise", desequilibre, uid=2)
        self.assertTrue(any("déséquilibrée" in e["message"] for e in r["rapport"]["errors"]))
        ok = classeur("tresorerie", "reprise", "REPRISE-TRESO-OK", {"Comptes_Tresorerie": [
            {"reference_externe": "CAISSE-USD", "libelle": "Caisse siège", "devise": "USD", "compte_comptable": "5711",
             "rattachement": "cash", "solde_ouverture": 250, "date_solde": D, "reference_releve": "PV comptage 31/10"},
            {"reference_externe": "BANQUE-CDF", "libelle": "Banque CDF", "devise": "CDF", "compte_comptable": "5212",
             "rattachement": "bank", "solde_ouverture": "1500000", "date_solde": D, "reference_releve": "Relevé 10/2026"}]})
        r = self.importer("tresorerie", "reprise", ok, uid=2)
        totaux = {(t["devise"], t["compte"]): t["montant"] for t in r["rapport"]["totaux"]}
        self.assertEqual(totaux[("USD", "CAISSE-USD")], 25000)
        self.assertEqual(totaux[("CDF", "BANQUE-CDF")], 150000000)          # pas de total multi-devises
        self.assertEqual(self.count("movements"), 0)
        self.activer(r["lot"]["id"])
        soldes = {a["ref"]: V.balance(a) for a in V.records("account")}
        self.assertEqual(soldes, {"CAISSE-USD": 25000, "BANQUE-CDF": 150000000})
        self.assertEqual(sum(x["debit"] - x["credit"] for x in V.ledger()), 0)
        rap = SV.lot(r["lot"]["id"])["rapprochement"]
        self.assertEqual({e["objet"] for e in rap["ecarts_activation"] if e["rubrique"] == "Trésorerie"}, {"CAISSE-USD", "BANQUE-CDF"})

    def test_remplacement_lot_prepare_et_rejet(self):
        self.budget_approuve()
        f1 = classeur("fournisseurs", "reprise", "REPRISE-R", {"Factures_Fournisseurs": [
            {"reference_externe": "FA9", "numero_facture": "F-9", "tiers_ref": "F001", "chantier_ref": "p",
             "date_facture": "2026-09-01", "date_echeance": "2026-12-01", "devise": "USD", "montant": 900}]})
        r1 = self.importer("fournisseurs", "reprise", f1, uid=2)
        f2 = classeur("fournisseurs", "reprise", "REPRISE-R", {"Factures_Fournisseurs": [
            {"reference_externe": "FA9", "numero_facture": "F-9", "tiers_ref": "F001", "chantier_ref": "p",
             "date_facture": "2026-09-01", "date_echeance": "2026-12-01", "devise": "USD", "montant": 950}]})
        self.as_(2)
        pv = SV.preview_import(SV.parse_file(f2, "fournisseurs"), SV.contexte("reprise", D), f2)
        with self.assertRaises(SV.ImportBloque):
            SV.commit_import(pv["preview_id"], {"accepter_alertes": True})
        r2 = SV.commit_import(pv["preview_id"], {"accepter_alertes": True, "remplacer": True})
        self.assertEqual(SV.lot(r1["lot"]["id"])["statut"], "Remplacé")
        statuts = sorted(c["data"]["status"] for c in V.records("carryover"))
        self.assertEqual(statuts, ["Rejetée", "À approuver"])
        self.activer(r2["lot"]["id"])
        self.as_(3)
        with self.assertRaises(ValueError):
            SV.rejeter(r2["lot"]["id"], "trop tard")


class RemplacementBudget(Base):
    def test_budget_rejete_puis_lot_corrige(self):
        def f(montant):
            return classeur("budgets", "reprise", "REPRISE-B2", {
                "Budgets": [{"reference_externe": "CT-B2", "chantier_ref": "p", "devise": "USD", "montant_contrat": 100,
                             "cout_estime_terminaison": montant}],
                "Lignes_Budget": [{"budget_ref": "CT-B2", "numero_ligne": 1, "code_dqe": "MAT", "libelle": "Matériaux",
                                   "budget_initial": montant}]})
        r1 = self.importer("budgets", "reprise", f(80), uid=2)
        self.as_(3)
        SV.rejeter(r1["lot"]["id"], "Montant erroné")
        self.assertEqual(V.records("budget")[0]["data"]["status"], "Rejeté")
        r2 = self.importer("budgets", "reprise", f(90), uid=2)
        self.assertEqual(r2["lot"]["statut"], "Préparé")
        self.activer(r2["lot"]["id"])
        actifs = V.active_budgets(V.records("budget"))
        self.assertEqual([(b["ref"], b["data"]["status"], b["data"]["initial"]) for b in actifs], [("CT-B2", "Approuvé", 9000)])
        self.as_(1)
        self.assertRaises(ValueError, V.guard_legacy, "p")


class Workflow(Base):
    def test_besoins_en_brouillon_sans_validation_importee(self):
        self.budget_approuve(10000000)
        contenu = classeur("besoins", "en_cours", "BESOINS-OCT", {"Besoins": [
            {"reference_externe": "DA-001", "chantier_ref": "Résidence Kintambo", "code_dqe": "MAT", "objet": "Ciment 50 sacs",
             "beneficiaire_ref": "F001", "nature": "Achat", "quantite": 50, "prix_unitaire": "12,50", "devise": "USD",
             "echeance": "2026-11-15"}]})
        r = self.importer("besoins", "en_cours", contenu, uid=1)
        op = V.records("operation")[0]
        self.assertEqual((op["data"]["status"], op["data"]["level"], op["data"]["approvals"], op["data"]["amount"],
                          op["data"]["creator"]), ("Brouillon", 0, [], 62500, 1))
        self.as_(1)
        V.submit(op["id"])
        with self.assertRaises(ValueError):
            V.approve(op["id"], "auto")
        self.as_(8)
        V.approve(op["id"], "N1 contrôlé")
        self.assertEqual(db.un("v7_records", op["id"])["data"]["level"], 1)

    def test_besoin_sans_budget_approuve_refuse_et_devise(self):
        contenu = classeur("besoins", "en_cours", "BESOINS-KO", {"Besoins": [
            {"reference_externe": "DA-002", "chantier_ref": "p", "code_dqe": "MAT", "objet": "Sable", "beneficiaire_ref": "F001",
             "nature": "Dépense", "montant": 100, "devise": "CDF", "echeance": "2026-11-15"}]})
        r = self.importer("besoins", "en_cours", contenu, uid=1)
        self.assertTrue(r["rapport"]["errors"])
        self.assertEqual(V.records("operation"), [])

    def test_activation_dg_distincte_et_finance_requise(self):
        contenu = classeur("budgets", "reprise", "REPRISE-WF", {
            "Budgets": [{"reference_externe": "CT-WF", "chantier_ref": "p", "devise": "USD", "montant_contrat": 100,
                         "cout_estime_terminaison": 100}],
            "Lignes_Budget": [{"budget_ref": "CT-WF", "numero_ligne": 1, "code_dqe": "MAT", "libelle": "Matériaux",
                               "budget_initial": 100}]})
        r = self.importer("budgets", "reprise", contenu, uid=4)          # importé par une DG
        self.as_(5)
        with self.assertRaises(ValueError):
            SV.activer(r["lot"]["id"], "sans Finance")
        self.as_(3)
        SV.valider_finance(r["lot"]["id"], "ok")
        self.as_(4)
        with self.assertRaises(ValueError):
            SV.activer(r["lot"]["id"], "DG importatrice")
        self.as_(3)
        with self.assertRaises(ValueError):
            SV.activer(r["lot"]["id"], "pas DG")
        self.as_(5)
        SV.activer(r["lot"]["id"], "ok")
        self.assertEqual(V.records("budget")[0]["data"]["status"], "Approuvé")


@unittest.skipUnless(PG, "Concurrence réelle testée sur PostgreSQL (BG_TEST_PG)")
class Concurrence(Base):
    def test_deux_confirmations_simultanees(self):
        contenu = classeur("tiers", "en_cours", "LOT-CONC", {"Tiers": [
            {"reference_externe": f"T{i}", "nom": f"Tiers {i}", "roles": "client"} for i in range(30)]})
        self.as_(2)
        ds = SV.parse_file(contenu, "tiers")
        apercus = [SV.preview_import(ds, SV.contexte("en_cours"), contenu)["preview_id"] for _ in range(4)]
        barriere, resultats = threading.Barrier(4), []
        def go(pid):
            barriere.wait()
            try:
                resultats.append(SV.commit_import(pid, {})["existant"])
            except Exception as e:  # noqa: BLE001
                resultats.append(type(e).__name__)
        fils = [threading.Thread(target=go, args=(p,)) for p in apercus]
        for f in fils:
            f.start()
        for f in fils:
            f.join()
        self.assertEqual(resultats.count(False), 1, resultats)
        self.assertEqual((self.count("import_lots"), self.count("clients")), (1, 31))

    def test_deux_lots_meme_facture_simultanes(self):
        self.budget_approuve()
        def f(ref):
            return classeur("fournisseurs", "reprise", ref, {"Factures_Fournisseurs": [
                {"reference_externe": "FA-CONC", "numero_facture": "F-1", "tiers_ref": "F001", "chantier_ref": "p",
                 "date_facture": "2026-09-01", "date_echeance": "2026-12-01", "devise": "USD", "montant": 100}]})
        self.as_(2)
        pids = []
        for ref in ("LOT-A", "LOT-B"):
            c = f(ref)
            pids.append(SV.preview_import(SV.parse_file(c, "fournisseurs"), SV.contexte("reprise", D), c)["preview_id"])
        barriere, resultats = threading.Barrier(2), []
        def go(pid):
            barriere.wait()
            try:
                resultats.append(SV.commit_import(pid, {"accepter_alertes": True})["lot"]["statut"])
            except Exception as e:  # noqa: BLE001
                resultats.append(type(e).__name__)
        fils = [threading.Thread(target=go, args=(p,)) for p in pids]
        for x in fils:
            x.start()
        for x in fils:
            x.join()
        self.assertEqual(len([c for c in V.records("carryover") if c["data"]["status"] != "Rejetée"]), 1, resultats)


class Sauvegarde(Base):
    def test_sauvegarde_restauration_lots_correspondances_pieces_parametres(self):
        self.as_(4)
        with db.transaction("limites") as t:
            t.parametre("imports_limites", {"lignes": 3000})
        contenu = classeur("globale", "en_cours", "LOT-SAV", {
            "Tiers": [{"reference_externe": "C1", "nom": "Client Un", "roles": "client"}],
            "Documents": [{"reference_externe": "DOC1", "entite_cible": "tiers", "reference_cible": "C1",
                           "nom_fichier": "contrat.pdf", "type": "Document de chantier", "date": "2026-10-01"}]})
        r = self.importer("globale", "en_cours", contenu, uid=4, fichiers={"contrat.pdf": b"%PDF-contrat-test"})
        self.assertEqual(r["lot"]["statut"], "Importé")
        self.assertEqual(len([p for p in db.tout("pieces") if p["batch_id"] == r["lot"]["id"]]), 1)
        sauvegarde = migration.export_json()
        migration.import_json(sauvegarde)
        self.assertEqual(self.count("import_lots"), 1)
        self.assertEqual({m["reference_externe"] for m in db.tout("import_mappings")}, {"C1", "DOC1"})
        self.assertEqual(self.count("pieces"), 1)
        self.assertEqual(db.parametre("imports_limites"), {"lignes": 3000})
        self.assertEqual(X.limites()["lignes"], 3000)
        self.assertIsNotNone(SV.contenu_lot(r["lot"]["id"]))


if __name__ == "__main__":
    unittest.main()
