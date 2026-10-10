"""Lecture bornée des classeurs, refus des formules et macros, normalisation et génération des modèles.

Règle des nombres (explicite) : un seul séparateur décimal, virgule OU point, sans séparateur de milliers.
« 1234,56 » et « 1234.56 » sont acceptés ; « 1 234,56 », « 1.234,56 » ou « 1,234.56 » sont refusés comme ambigus.
Montants : deux décimales au plus, calculés avec Decimal puis stockés en centimes entiers.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from imports import schema as S

LIMITES_DEFAUT = {"taille_fichier": 8 * 1024 * 1024, "lignes": 5000, "cellules": 400_000,
                  "decompresse": 64 * 1024 * 1024, "ratio": 120, "entrees_zip": 400}
MONTANT_MAX = Decimal(10) ** 12          # 1 000 milliards d'unités, soit 10^14 centimes (limite V7)
QTE_MAX = Decimal(10) ** 9


def limites() -> dict:
    try:
        import db
        conf = db.parametre("imports_limites") or {}
    except Exception:  # noqa: BLE001 - hors application (tests unitaires sans base)
        conf = {}
    return {**LIMITES_DEFAUT, **{k: int(v) for k, v in conf.items() if k in LIMITES_DEFAUT}}


def empreinte(contenu: bytes) -> str:
    return hashlib.sha256(contenu).hexdigest()


# ------------------------------------------------------------------ contrôle de l'archive .xlsx
def _controle_zip(contenu: bytes, lim: dict) -> None:
    if len(contenu) > lim["taille_fichier"]:
        raise ValueError(f"Fichier trop volumineux : {lim['taille_fichier'] // (1024 * 1024)} Mo au plus.")
    if not contenu.startswith(b"PK"):
        raise ValueError("Format refusé : seul le format .xlsx (Excel 2007 et suivants) est accepté.")
    try:
        z = zipfile.ZipFile(io.BytesIO(contenu))
    except zipfile.BadZipFile as e:
        raise ValueError("Fichier .xlsx illisible ou corrompu.") from e
    infos = z.infolist()
    if len(infos) > lim["entrees_zip"]:
        raise ValueError("Classeur anormalement complexe (trop d’éléments internes).")
    total = sum(i.file_size for i in infos)
    if total > lim["decompresse"] or total > max(1, len(contenu)) * lim["ratio"]:
        raise ValueError("Taux de décompression anormal : fichier refusé par sécurité.")
    noms = {i.filename.lower() for i in infos}
    if any(n.endswith("vbaproject.bin") or n.startswith("xl/activex") or n.endswith(".bin") for n in noms):
        raise ValueError("Macros ou objets actifs détectés : enregistrez le classeur au format .xlsx sans macros.")
    try:
        types = z.read("[Content_Types].xml").decode("utf-8", "ignore")
    except KeyError as e:
        raise ValueError("Fichier .xlsx incomplet.") from e
    if "macroEnabled" in types or "vbaProject" in types:
        raise ValueError("Classeur prenant en charge les macros (.xlsm) refusé.")
    if any(n.startswith("xl/externallinks/") for n in noms):
        raise ValueError("Liaisons vers d’autres classeurs détectées : remplacez-les par des valeurs.")


# ------------------------------------------------------------------ normalisation des valeurs
def texte(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "oui" if v else "non"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            raise ValueError("valeur non finie")
        if v.is_integer():
            return str(int(v))
        return repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v).strip()


def decimal_strict(v, libelle: str = "Nombre") -> Decimal:
    if isinstance(v, bool):
        raise ValueError(f"{libelle} : nombre attendu.")
    if isinstance(v, (int, float)):
        d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
    else:
        raw = str(v).strip().replace(" ", " ")
        if not re.fullmatch(r"-?\d+(?:[.,]\d+)?", raw):
            raise ValueError(f"{libelle} : nombre ambigu ou invalide « {raw} » (un seul séparateur décimal, "
                             "sans séparateur de milliers).")
        d = Decimal(raw.replace(",", "."))
    if not d.is_finite():
        raise ValueError(f"{libelle} : valeur non finie.")
    return d


def centimes(v, libelle: str = "Montant", negatif: bool = False) -> int:
    d = decimal_strict(v, libelle)
    if d < 0 and not negatif:
        raise ValueError(f"{libelle} : montant négatif refusé.")
    if abs(d) > MONTANT_MAX:
        raise ValueError(f"{libelle} : montant hors limites.")
    c = d * 100
    if c != c.to_integral_value():
        raise ValueError(f"{libelle} : deux décimales au maximum.")
    return int(c)


def normaliser(champ: S.Champ, v):
    """Renvoie la valeur normalisée, ou lève ValueError avec un message lisible."""
    vide = v is None or (isinstance(v, str) and v.strip() == "")
    if vide:
        if champ.requis:
            raise ValueError("Champ obligatoire manquant.")
        return None
    t = champ.type
    if t in ("money",):
        return centimes(v, "Montant")
    if t in ("qty", "pct", "decimal"):
        d = decimal_strict(v, "Nombre")
        if d < 0:
            raise ValueError("Nombre négatif refusé.")
        if t == "pct" and d > 100:
            raise ValueError("Pourcentage supérieur à 100.")
        if d > QTE_MAX:
            raise ValueError("Nombre hors limites.")
        if t == "decimal" and d == 0:
            raise ValueError("Valeur nulle refusée.")
        return str(d.normalize()) if t == "decimal" else float(d)
    if t == "int":
        d = decimal_strict(v, "Entier")
        if d != d.to_integral_value() or d < 0 or d > 10**6:
            raise ValueError("Entier positif attendu.")
        return int(d)
    if t == "date":
        if isinstance(v, dt.datetime):
            if v.time() != dt.time():
                raise ValueError("Date sans heure attendue (AAAA-MM-JJ).")
            v = v.date()
        if isinstance(v, dt.date):
            return v.isoformat()
        s = str(v).strip()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            raise ValueError("Date au format AAAA-MM-JJ attendue (les formats JJ/MM/AAAA sont ambigus).")
        try:
            dt.date.fromisoformat(s)
        except ValueError as e:
            raise ValueError("Date inexistante.") from e
        return s
    if t == "month":
        if isinstance(v, (dt.date, dt.datetime)):
            raise ValueError("Période AAAA-MM attendue en texte.")
        s = str(v).strip()
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", s):
            raise ValueError("Période au format AAAA-MM attendue.")
        return s
    if t == "bool":
        s = texte(v).lower()
        if s in ("oui", "1", "vrai", "true", "o"):
            return True
        if s in ("non", "0", "faux", "false", "n"):
            return False
        raise ValueError("Oui ou non attendu.")
    s = texte(v)
    if t == "ref":
        if isinstance(v, float) and not v.is_integer():
            raise ValueError("Référence en texte attendue (format de cellule Texte).")
        if not re.fullmatch(r"[\w.\-/: ]{1,120}", s, flags=re.UNICODE):
            raise ValueError("Référence : lettres, chiffres, espace, . - / : _ (120 caractères au plus).")
        return s
    if t == "enum":
        if s not in champ.valeurs:
            raise ValueError("Valeur non autorisée : " + ", ".join(map(str, champ.valeurs)) + ".")
        return s
    if t == "filename":
        if "/" in s or "\\" in s or s.startswith(".") or ".." in s or ":" in s:
            raise ValueError("Nom de fichier sans chemin attendu.")
        if not re.fullmatch(r"[\w .\-()]{1,200}\.(pdf|jpg|jpeg|png)", s, flags=re.IGNORECASE | re.UNICODE):
            raise ValueError("Nom de fichier PDF, JPEG ou PNG attendu.")
        return s
    if t == "sha256":
        s = s.lower()
        if not re.fullmatch(r"[0-9a-f]{64}", s):
            raise ValueError("Empreinte SHA-256 de 64 caractères hexadécimaux attendue.")
        return s
    if len(s) > champ.max_len:
        raise ValueError(f"Texte trop long ({champ.max_len} caractères au plus).")
    if s[:1] in ("=", "+", "@") or (s[:1] == "-" and len(s) > 1 and not s[1:2].isdigit()):
        raise ValueError("Texte commençant par un caractère de formule refusé.")
    return s


# ------------------------------------------------------------------ lecture du classeur
def lire(contenu: bytes, feuilles_permises: tuple) -> dict:
    """Lecture bornée. Renvoie {params, feuilles: {nom: [lignes brutes]}, errors, empreinte}."""
    lim = limites()
    _controle_zip(contenu, lim)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), data_only=False, read_only=False, keep_links=False)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Classeur illisible : utilisez le modèle .xlsx fourni.") from e
    errors: list[dict] = []
    def err(feuille, ligne, champ, msg):
        errors.append({"sheet": feuille, "line": ligne, "field": champ, "message": msg})
    params, feuilles, total_lignes, total_cellules = {}, {}, 0, 0
    noms_vus = set()
    for ws in wb.worksheets:
        nom = ws.title
        if nom in noms_vus:
            err(nom, 1, "", "Feuille en double.")
            continue
        noms_vus.add(nom)
        if nom == "Instructions":
            continue
        total_cellules += (ws.max_row or 0) * (ws.max_column or 0)
        if total_cellules > lim["cellules"]:
            raise ValueError("Classeur trop grand : nombre de cellules supérieur à la limite autorisée.")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    err(nom, cell.row, cell.coordinate, "Formule refusée : remplacez-la par sa valeur.")
        if nom == "Parametres":
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r and r[0] not in (None, ""):
                    cle = texte(r[0])
                    if cle in params:
                        err(nom, 0, cle, "Paramètre en double.")
                    params[cle] = r[1] if len(r) > 1 else None
            continue
        if nom not in S.FEUILLES:
            err(nom, 1, "", "Feuille inconnue : utilisez le modèle fourni sans ajouter de feuille.")
            continue
        if nom not in feuilles_permises:
            err(nom, 1, "", "Feuille non prévue pour cette fonction d’import.")
            continue
        for plage in ws.merged_cells.ranges:
            err(nom, plage.min_row, str(plage), "Cellules fusionnées refusées dans la zone de données.")
        grille = list(ws.iter_rows(values_only=True))
        entetes = [texte(v) for v in (grille[0] if grille else [])]
        while entetes and entetes[-1] == "":
            entetes.pop()
        attendu = S.FEUILLES[nom].colonnes
        if len(set(entetes)) != len(entetes):
            err(nom, 1, "", "En-têtes répétés.")
        if any("approuv" in h.lower() or h.lower().startswith("statut_valid") for h in entetes):
            err(nom, 1, "", "Un statut « approuvé » ne peut pas être importé : la validation se fait dans l’application.")
        if entetes != attendu:
            manquants = [c for c in attendu if c not in entetes]
            en_trop = [c for c in entetes if c not in attendu]
            err(nom, 1, "", "Colonnes différentes du modèle " + S.VERSION_MODELE
                + (f" ; manquantes : {', '.join(manquants)}" if manquants else "")
                + (f" ; inconnues : {', '.join(en_trop)}" if en_trop else "")
                + ("" if manquants or en_trop else " ; ordre modifié") + ".")
            continue
        lignes = []
        for i, vals in enumerate(grille[1:], start=2):
            if all(v in (None, "") for v in vals):
                continue
            if any(v not in (None, "") for v in vals[len(entetes):]):
                err(nom, i, "", "Valeurs hors des colonnes du modèle.")
            lignes.append({"__line": i, **{h: (vals[j] if j < len(vals) else None) for j, h in enumerate(entetes)}})
        total_lignes += len(lignes)
        feuilles[nom] = lignes
    if total_lignes > lim["lignes"]:
        err("Classeur", 0, "", f"{total_lignes} lignes utiles : {lim['lignes']} au plus par lot. "
                               "Découpez en plusieurs lots indépendants.")
    if "Parametres" not in noms_vus:
        err("Parametres", 0, "", "Feuille Parametres absente : utilisez le modèle fourni.")
    return {"params": params, "feuilles": feuilles, "errors": errors, "empreinte": empreinte(contenu),
            "taille": len(contenu)}


# ------------------------------------------------------------------ modèles téléchargeables
GRAS = Font(bold=True, color="FFFFFF")
FOND = PatternFill("solid", fgColor="112740")
FOND_REQ = PatternFill("solid", fgColor="8A3B12")


def modele(cle_fonction: str, mode: str = "", date_bascule: str = "", reference_lot: str = "") -> bytes:
    f = S.fonction(cle_fonction)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Instructions")
    lignes = [
        [f"MY DESTINY SARL — Modèle d’import « {f.libelle} » (version {S.VERSION_MODELE})"],
        [f.description],
        [""],
        ["1. Renseignez la feuille Parametres : mode_import (archives, reprise ou en_cours), date_bascule "
         "(AAAA-MM-JJ, obligatoire en reprise) et reference_lot (unique)."],
        ["2. Ne modifiez ni les noms des feuilles ni les en-têtes. Laissez vides les feuilles non utilisées. "
         "Pas de formules, de macros ni de cellules fusionnées."],
        ["3. Dates AAAA-MM-JJ ; périodes AAAA-MM ; références en texte (les zéros initiaux sont conservés)."],
        ["4. Nombres : un seul séparateur décimal (virgule ou point), sans séparateur de milliers ; "
         "montants à deux décimales au plus. Devises USD, CDF, EUR tenues séparément, sans conversion."],
        ["5. Les références (client_ref, chantier_ref, matricule…) renvoient à une ligne du même classeur ou à "
         "une fiche existante. Le serveur résout les identifiants internes."],
        ["6. Aucune colonne « statut approuvé » : budgets, soldes et reliquats sont approuvés dans l’application "
         "par les personnes habilitées. L’import ne crée ni signature ni validation."],
        ["7. Reprise : les paiements antérieurs restent en archives ; seul le reliquat est repris, sans nouveau "
         "décaissement. Une justification déclarée doit être prouvée par les pièces."],
        ["8. Les pièces (PDF, JPEG, PNG) se chargent séparément avec la feuille Documents comme manifeste."],
        [""], ["Feuille", "Champ", "Obligatoire", "Type", "Valeurs / aide"],
    ]
    for l in lignes:
        ws.append(l)
    for nom in f.feuilles:
        fe = S.FEUILLES[nom]
        etat = "" if fe.statut == "réalisé" else " (archives seulement)"
        for c in fe.champs:
            ws.append([nom + etat, c.nom, "oui" if c.requis else "", c.type,
                       (", ".join(map(str, c.valeurs)) if c.valeurs else c.aide)
                       + (" — confidentiel RH/DG" if c.confidentiel else "")])
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 28 if col != "E" else 70
    p = wb.create_sheet("Parametres")
    p.append(["parametre", "valeur"])
    for k, v in [("version_modele", S.VERSION_MODELE), ("source_systeme", ""), ("mode_import", mode),
                 ("date_bascule", date_bascule), ("reference_lot", reference_lot)]:
        p.append([k, v])
    for c in p[1]:
        c.font, c.fill = GRAS, FOND
    p.column_dimensions["A"].width = 22
    p.column_dimensions["B"].width = 40
    for row in p.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = "@"
    dv = DataValidation(type="list", formula1='"archives,reprise,en_cours"', allow_blank=False)
    p.add_data_validation(dv)
    dv.add("B4")
    for nom in f.feuilles:
        fe = S.FEUILLES[nom]
        ws = wb.create_sheet(nom)
        ws.append(fe.colonnes)
        for j, c in enumerate(fe.champs, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font, cell.fill = GRAS, (FOND_REQ if c.requis else FOND)
            cell.alignment = Alignment(wrap_text=True)
            lettre = cell.column_letter
            ws.column_dimensions[lettre].width = max(16, len(c.nom) + 3)
            if c.type in ("ref", "text", "date", "month", "filename", "sha256", "enum"):
                for r in range(2, 502):
                    ws.cell(row=r, column=j).number_format = "@"
            if c.type == "enum" and len(",".join(map(str, c.valeurs))) < 240:
                v = DataValidation(type="list", formula1='"' + ",".join(map(str, c.valeurs)) + '"', allow_blank=True)
                ws.add_data_validation(v)
                v.add(f"{lettre}2:{lettre}501")
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rapport(report: dict) -> bytes:
    """Rapport de contrôle téléchargeable (erreurs, alertes, effets, totaux par devise)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle"
    ws.append(["Niveau", "Feuille", "Ligne", "Champ", "Message"])
    for e in report.get("errors", []):
        ws.append(["Erreur bloquante", e["sheet"], e["line"], e["field"], e["message"]])
    for w in report.get("warnings", []):
        ws.append(["Alerte" + (" (décision requise)" if w.get("decision") else ""), w["sheet"], w["line"],
                   w["field"], w["message"]])
    if ws.max_row == 1:
        ws.append(["Information", "", "", "", "Aucune anomalie détectée."])
    ef = wb.create_sheet("Effets")
    ef.append(["Ordre", "Feuille", "Ligne", "Action", "Référence", "Détail"])
    for a in report.get("plan", []):
        ef.append([a.get("ordre"), a["sheet"], a["line"], a["action"], a.get("ref", ""), a.get("detail", "")])
    tt = wb.create_sheet("Totaux")
    tt.append(["Rubrique", "Devise", "Compte", "Montant"])
    for t in report.get("totaux", []):
        tt.append([t["rubrique"], t["devise"], t.get("compte", ""), t["montant"] / 100])
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font, c.fill = GRAS, FOND
        sheet.column_dimensions["E"].width = 90
        sheet.column_dimensions["F"].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
